"""核查 pass_ratio 真伪: 真实报告的检测方案表列头是什么?

线索: vision_cache 里 pass_ratio 分布极可疑
  长兴 {3/1: 89, 1/1: 15}
  新泰 {3/1: 17, 0.5/5: 5, 2/5: 4, ...}
而我的 prompt 里给的示例正是 "如3/1" → 强烈怀疑**模型照抄示例**。

真实报告(预评备案稿)里若引用/复制了检测方案表, 就有**文本层**可读。
搜: "相对稳定" / "检测点数" / "采样点数" / "检测方案"
"""
import re
from pathlib import Path
from docx import Document

DOCS = {
    "新泰备案稿": "/tmp/xt_probe/6-1新泰--预评（备案稿）.docx",
}

print("=" * 80)
print("① 搜 '相对稳定' 的上下文")
print("=" * 80)
for tag, p in DOCS.items():
    if not Path(p).exists():
        print(f"  [{tag}] 不存在: {p}")
        continue
    doc = Document(p)
    paras = [x.text.strip() for x in doc.paragraphs if x.text.strip()]
    for i, t in enumerate(paras):
        if "相对稳定" in t or "相对不稳定" in t:
            print(f"[{tag}][段{i}] {t[:250]}")

print()
print("=" * 80)
print("② 表格里搜: 含'原辅料'或'相对稳定'或'检测点数'的表")
print("=" * 80)
for tag, p in DOCS.items():
    if not Path(p).exists():
        continue
    doc = Document(p)
    for ti, tb in enumerate(doc.tables):
        txt = "\n".join(c.text for r in tb.rows for c in r.cells)
        if ("原辅料" in txt) or ("相对稳定" in txt) or ("检测点数" in txt) or ("采样点数" in txt):
            print(f"\n--- [{tag}] 表{ti} ({len(tb.rows)}行 x {len(tb.columns)}列) ---")
            for ri, r in enumerate(tb.rows[:8]):
                cells = [c.text.strip().replace("\n", "/") for c in r.cells]
                print(f"   r{ri}: {' | '.join(cells)[:250]}")
            if len(tb.rows) > 8:
                print(f"   ... 共 {len(tb.rows)} 行")

print()
print("=" * 80)
print("③ 搜 分母为5的分数样值 (\\d/5) 与 3/1")
print("=" * 80)
for tag, p in DOCS.items():
    if not Path(p).exists():
        continue
    doc = Document(p)
    full = "\n".join([x.text for x in doc.paragraphs]
                     + [c.text for tb in doc.tables for r in tb.rows for c in r.cells])
    for pat in [r"\d/5\b", r"\b3/1\b", r"\d/\d"]:
        hits = re.findall(pat, full)
        print(f"  [{tag}] {pat!r}: {len(hits)} 次  样本={hits[:10]}")

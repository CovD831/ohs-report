"""原报告 (长兴备案稿) 结构画像 — 供生成报告逐项对照

输出: 章节树 / 每章字数 / 表格数 / 表题清单
"""
import json
import re
import sys
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

SRC = Path(sys.argv[1] if len(sys.argv) > 1 else
           "/tmp/bench_orig/长兴--预评（备案稿7-31）.docx")


def iter_blocks(doc):
    """按文档顺序产出段落/表格"""
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def cap_of(t: str) -> bool:
    return bool(re.match(r"^表\s*[\d.\-—]+", (t or "").strip()))


def main():
    doc = docx.Document(str(SRC))
    heads = []          # (级别, 文本)
    paras = 0
    chars = 0
    caps = []
    tbl_rows = 0
    cur = None
    per_sec = {}        # 章 -> [字数, 表数]

    for b in iter_blocks(doc):
        if isinstance(b, Paragraph):
            t = b.text.strip()
            if not t:
                continue
            st = (getattr(b.style, "name", "") or "")
            if st.startswith("Heading"):
                m = re.match(r"^(\d+(?:\.\d+){0,3})\s*(.*)", t)
                if m:
                    lvl = m.group(1).count(".") + 1
                    heads.append((lvl, m.group(1), m.group(2)[:40]))
                    cur = m.group(1)
                    per_sec.setdefault(cur, [0, 0])
            if cap_of(t):
                caps.append(t[:60])
                if cur:
                    per_sec.setdefault(cur, [0, 0])[1] += 1
            paras += 1
            chars += len(t)
            if cur:
                per_sec.setdefault(cur, [0, 0])[0] += len(t)
        else:
            tbl_rows += len(b.rows)

    print("=" * 78)
    print(f"原报告: {SRC.name}")
    print(f"  段落 {paras} | 字符 {chars} | 表格 {len(doc.tables)} | 表格行 {tbl_rows} | 表题 {len(caps)}")
    print()
    print("=== 章节标题 (前 60) ===")
    for lvl, num, txt in heads[:60]:
        print(f"  {'  ' * (lvl - 1)}{num} {txt}")
    print()
    print(f"共 {len(heads)} 个标题")
    out = {"paras": paras, "chars": chars, "tables": len(doc.tables),
           "table_rows": tbl_rows, "captions": caps, "heads": heads,
           "per_sec": per_sec}
    Path("/tmp/orig_profile.json").write_text(
        json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    print("→ /tmp/orig_profile.json")


if __name__ == "__main__":
    main()

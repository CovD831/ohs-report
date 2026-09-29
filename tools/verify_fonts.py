"""验证标题/表题/表格字体已改为仿宋 (对齐原报告)"""
import re
import sys
from collections import Counter

import docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

GEN = "/tmp/report_fixed.docx"
ORIG = "/tmp/bench_orig/长兴--预评（备案稿7-31）.docx"


def scan_headings(path):
    d = docx.Document(path)
    c = Counter()
    for ch in d.element.body.iterchildren():
        if ch.tag != qn("w:p"):
            continue
        x = Paragraph(ch, d)
        t = x.text.strip()
        if not t or len(t) > 46:
            continue
        m = re.match(r"^(\d{1,2}(?:\.\d{1,2}){0,3})[\s、.．]{1,3}(\S.*)$", t)
        if not m or m.group(2)[:1].isdigit():
            continue
        lvl = m.group(1).count(".") + 1
        if lvl > 4:
            continue
        r = x.runs[0] if x.runs else None
        f = sz = None
        if r is not None and r._element.rPr is not None:
            ft = r._element.rPr.find(qn("w:rFonts"))
            if ft is not None:
                f = ft.get(qn("w:eastAsia"))
            s = r._element.rPr.find(qn("w:sz"))
            sz = s.get(qn("w:val")) if s is not None else None
        c[(lvl, f, sz)] += 1
    return c


def scan_captions(path):
    d = docx.Document(path)
    c = Counter()
    for ch in d.element.body.iterchildren():
        if ch.tag != qn("w:p"):
            continue
        x = Paragraph(ch, d)
        t = x.text.strip()
        if not re.match(r"^表\s*[\d.]+[-—]", t):
            continue
        r = x.runs[0] if x.runs else None
        f = sz = bold = None
        if r is not None and r._element.rPr is not None:
            ft = r._element.rPr.find(qn("w:rFonts"))
            if ft is not None:
                f = ft.get(qn("w:eastAsia"))
            s = r._element.rPr.find(qn("w:sz"))
            sz = s.get(qn("w:val")) if s is not None else None
            # ⚠ w:b 元素存在 ≠ 加粗: val="0" 表示**明确不加粗**。
            #   只判存在会把 bold=False 误报成 True (实测踩过)。
            bb = r._element.rPr.find(qn("w:b"))
            bold = bb is not None and bb.get(qn("w:val")) not in ("0", "false")
        c[(f, sz, bold)] += 1
    return c


def scan_cells(path):
    d = docx.Document(path)
    c = Counter()
    for t in d.tables[:14]:
        for row in t.rows[:6]:
            for cell in row.cells[:5]:
                for para in cell.paragraphs:
                    for r in para.runs:
                        # ⚠ 只统计**有文本**的 run: 空白 run 无字体属正常,
                        #   早期统计把空 run 计入 → 误报大量 None (实测 1701 空 vs 16 实质)
                        if not r.text.strip():
                            continue
                        f = sz = None
                        if r._element.rPr is not None:
                            ft = r._element.rPr.find(qn("w:rFonts"))
                            if ft is not None:
                                f = ft.get(qn("w:eastAsia"))
                            s = r._element.rPr.find(qn("w:sz"))
                            sz = s.get(qn("w:val")) if s is not None else None
                        c[(f, sz)] += 1
    return c


def show(label, c, n=4):
    print(f"  {label}:")
    for k, v in c.most_common(n):
        print(f"     {v:>5} 次  {k}")


print("=" * 74)
print("标题字体")
print("=" * 74)
show("原报告", scan_headings(ORIG))
show("生成(修正后)", scan_headings(GEN))
print()
print("=" * 74)
print("表题字体")
print("=" * 74)
show("原报告", scan_captions(ORIG), 3)
show("生成(修正后)", scan_captions(GEN), 3)
print()
print("=" * 74)
print("表格单元格字体")
print("=" * 74)
show("原报告", scan_cells(ORIG), 3)
show("生成(修正后)", scan_cells(GEN), 3)

#!/usr/bin/env python3
"""h11/h14 设计取证 — 紧凑输出 (两备案稿物理表形态 + 缓存行形态 + 本地库分布)"""
import json, re
from pathlib import Path
from collections import Counter
from docx import Document

REPO = Path('/Users/abaaba/Projects/ohs-report')
PHYS = ('噪声', '照度', '工频电场', '高温', 'WBGT', '紫外辐射', '手传振动', '全身振动', '微波辐射', '高频电磁场')

def t26(v):
    s = str(v)
    return s[:26] + ('…' if len(s) > 26 else '')

print('########## A. 备案稿 docx: 物理/检测相关表 ##########')
for label, path in (('长兴', '/tmp/bench_orig/长兴--预评（备案稿7-31）.docx'),
                    ('新泰', '/tmp/xt_probe/6-1新泰--预评（备案稿）.docx')):
    p = Path(path)
    if not p.exists():
        print(f'-- {label}: MISSING {path}')
        continue
    doc = Document(path)
    hits = 0
    for ti, t in enumerate(doc.tables):
        first2 = ' '.join(c.text for r in t.rows[:2] for c in r.cells)
        if re.search(r'LEX|dB|照度|lx|工频|噪声', first2):
            hits += 1
            print(f'== {label} 表{ti} {len(t.rows)}x{len(t.columns)}')
            for ri in range(min(5, len(t.rows))):
                cells = ' | '.join(c.text.replace('\n', '/').replace('\r', '')[:13] for c in t.rows[ri].cells)
                print('   r%d %s' % (ri, cells[:230]))
    print(f'-- {label}: 命中表 {hits}/{len(doc.tables)}')
    n = 0
    for par in doc.paragraphs:
        tx = par.text.strip()
        if 0 < len(tx) < 44 and re.match(r'^表\s*\d', tx) and re.search(r'检测|噪声|照度|结果|物理', tx):
            print('   图题:', tx)
            n += 1
            if n > 22:
                break

print()
print('########## B. 缓存条目形态 ##########')
for tag, rel in (('新泰', 'data/vision_cache/a66517be268181e6.json'),
                 ('长兴', 'data/vision_cache/32e58e26097c8e58.json')):
    f = REPO / rel
    if not f.exists():
        print(f'-- {tag}: MISSING {rel}')
        continue
    d = json.loads(f.read_text())['detections']
    cst = [e for e in d if str(e.get('cstel') or '').strip()]
    print(f'== {tag}: {len(d)}条; cstel非空 {len(cst)}')
    for e in cst[:5]:
        print('  C| factor=%s ctwa=%s cstel=%s judge=%s tt=%s p=%s sp=%s' % (
            t26(e.get('factor')), t26(e.get('ctwa')), t26(e.get('cstel')),
            t26(e.get('judgement')), e.get('table_type'), e.get('_page'), t26(e.get('sampling_point'))))
    ph = [e for e in d if any(str(e.get('factor') or '').strip() == pk for pk in PHYS)]
    print(f'  -- 纯物理factor行 {len(ph)}; 分布: {dict(Counter(str(e.get("factor")) for e in ph))}')
    phv = [e for e in ph if [x for x in (e.get('results') or []) if str(x).strip() not in ('', '—', '-', '/')]]
    print(f'  -- 其中results非空 {len(phv)}')
    for e in phv[:10]:
        print('  P| f=%s res=%s ctwa=%s cstel=%s j=%s tt=%s p=%s sp=%s' % (
            t26(e.get('factor')), e.get('results'), t26(e.get('ctwa')), t26(e.get('cstel')),
            t26(e.get('judgement')), e.get('table_type'), e.get('_page'), t26(e.get('sampling_point'))))
    print('  -- 物理行judgement分布:', dict(Counter(t26(e.get('judgement')) for e in phv)))
    mix = [e for e in d if re.search(r'噪声|照度|工频|高温|WBGT', str(e.get('factor') or '')) and str(e.get('factor') or '').strip() not in PHYS]
    print(f'  -- 混合/别名含物理词行 {len(mix)} 样例x6:')
    for e in mix[:6]:
        print('  M| f=%s res=%s j=%s p=%s tt=%s' % (t26(e.get('factor')), e.get('results'), t26(e.get('judgement')), e.get('_page'), e.get('table_type')))

print()
print('########## C. 本地库 detections 分布 ##########')
import sqlite3
conn = sqlite3.connect(str(REPO / 'data/ohs.db'))
tabs = [r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")]
pt = [t for t in tabs if 'project' in t.lower() or 'data' in t.lower()]
print('project-like tables:', pt)
for t in pt:
    cols = [c[1] for c in conn.execute(f'PRAGMA table_info({t})')]
    print(' ', t, '->', cols)

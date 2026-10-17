"""P0 批次指标采集: 4 项目导出逐指标 + 校验器在既有正文上的命中统计.

用法: .venv/bin/python tools/p0_metrics.py <tag>   # tag ∈ {pre_p0, post_p0}
产出: /tmp/p0_metrics_<tag>.json + data/report_<pid>_<tag>.docx
"""
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")

TAG = sys.argv[1] if len(sys.argv) > 1 else "pre_p0"
PIDS = ["ada54603a9", "c8c7ff0a4d", "2863a36451", "57ffa2b3b0"]

from docx import Document  # noqa: E402
from tools.export_e2e import build_project  # noqa: E402
from knowledge.oel import connect  # noqa: E402
from knowledge.paths import DATA_DIR  # noqa: E402  (数据目录唯一产地, 认 OHS_DATA_DIR)
from knowledge.project_assess import assess_project  # noqa: E402
from web.projects_db import get_project_data as _get_project_data  # noqa: E402
from web.validators import validate_section  # noqa: E402
from web.word_export import export_docx  # noqa: E402


def docx_metrics(path: Path) -> dict:
    d = Document(str(path))
    paras = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    body_chars = sum(len(t) for t in paras)
    n_tbl, n_rows, cell_chars, pend = 0, 0, 0, 0
    by_table: dict[str, int] = {}
    seen = set()
    for t in d.tables:
        n_tbl += 1
        n_rows += len(t.rows)
        cap = ""
        prev = t._tbl.getprevious()
        for _ in range(4):
            if prev is None:
                break
            txt = "".join(n.text or "" for n in prev.iter() if n.tag.endswith("}t")).strip()
            if txt:
                cap = txt[:40]
                break
            prev = prev.getprevious()
        # ⚠ 不做合并单元格去重: lxml 元素代理对象的 id() 会跨访问复用, 去重会误吞计数。
        #   口径与既有基线 (capfix4=125 / optfix=117) 保持一致: 逐 cell 出现即计。
        for r in t.rows:
            for c in r.cells:
                s = c.text.strip()
                cell_chars += len(s)
                if "待补充" in s:
                    pend += 1
                    by_table[cap or "(无表题)"] = by_table.get(cap or "(无表题)", 0) + s.count("待补充")
    pend += sum(t.count("待补充") for t in paras)
    return {"bytes": path.stat().st_size, "tables": n_tbl, "table_rows": n_rows,
            "body_chars": body_chars, "table_chars": cell_chars,
            "nonempty_paras": len(paras), "pending_cells": pend,
            "pending_by_table": by_table}


def validator_metrics(pid: str, data: dict, assess: dict) -> dict:
    states = data.get("section_states") or {}
    per_rule: dict[str, int] = {}
    flagged, total = 0, 0
    for sec, st in states.items():
        text = (st or {}).get("text") or ""
        if len(text) < 30:
            continue
        total += 1
        try:
            v = validate_section(sec, text, data, assess)
        except Exception as e:  # 校验器本身抛错 = 必须暴露
            per_rule["EXCEPTION"] = per_rule.get("EXCEPTION", 0) + 1
            print(f"    !! validate_section({sec}) 抛错: {e}")
            continue
        issues = v.get("issues") or []
        if issues:
            flagged += 1
        for it in issues:
            rid = it.get("rule_id", "?")
            per_rule[rid] = per_rule.get(rid, 0) + 1
    return {"sections_checked": total, "sections_with_issues": flagged, "by_rule": per_rule}


out = {"tag": TAG, "projects": {}}
conn = connect()
for pid in PIDS:
    data = _get_project_data(pid)
    project = build_project(pid)
    assess = assess_project(conn, project)
    assess["_project_data"] = dict(project)
    docx_path = DATA_DIR / f"report_{pid}_{TAG}.docx"
    export_docx(project, assess, docx_path, section_states=data.get("section_states"))
    m = docx_metrics(docx_path)
    m["validators"] = validator_metrics(pid, data, assess)
    out["projects"][pid] = m
    print(f"{pid}: 表{m['tables']} 行{m['table_rows']} 待补充{m['pending_cells']} "
          f"| 校验 {m['validators']['sections_with_issues']}/{m['validators']['sections_checked']} "
          f"节有issue {m['validators']['by_rule']}")
conn.close()

agg = {"tables": 0, "table_rows": 0, "body_chars": 0, "table_chars": 0,
       "nonempty_paras": 0, "pending_cells": 0, "bytes": 0}
vr = {}
tot = fl = 0
for p in out["projects"].values():
    for k in agg:
        agg[k] += p[k]
    tot += p["validators"]["sections_checked"]
    fl += p["validators"]["sections_with_issues"]
    for k, v in p["validators"]["by_rule"].items():
        vr[k] = vr.get(k, 0) + v
out["aggregate"] = agg
out["validators_agg"] = {"sections_checked": tot, "sections_with_issues": fl, "by_rule": vr}
print("\n=== 汇总 ===")
print(json.dumps(out["aggregate"], ensure_ascii=False))
print(json.dumps(out["validators_agg"], ensure_ascii=False))
Path(f"/tmp/p0_metrics_{TAG}.json").write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
print(f"→ /tmp/p0_metrics_{TAG}.json")

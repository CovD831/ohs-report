"""完整导出验收: 36 张骨架表里, 导出的 docx 实际出现哪些表

关卡: 不看代码猜测, 直接打开导出的 docx 数表格 + 抓表题,
      与骨架名单对账 → 得出"哪些表真产出了 / 哪些丢了"。
"""
import json
import re
import sqlite3
import sys
from pathlib import Path

ROOT = next((p for p in (Path("/app"), Path(__file__).resolve().parent.parent)
             if (p / "web" / "app.py").exists()), Path(__file__).resolve().parent.parent)
sys.path.insert(0, str(ROOT))

from knowledge.oel import connect_ro  # noqa: E402  (只读工厂: 路径走 knowledge.paths 单一产地)

import docx  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

# 骨架表名 → 期望表题子串 —— **从产品真源推导, 不再手工维护别名表**。
# 病根复盘 (2026-10): 旧 CAP_HINT 是第三份手抄名单, 「设备明细表」hint 写 09-28 的
# 旧措辞「设备设施」, 真表题 10-03 定稿「本项目涉及的主要生产装置及设备」→ 两套真源
# 漂移, 表明明在导出里却被判「导出未见」(判据错, 不是产品缺陷)。
# 现在期望表题直接推自 _SUB_TABLE_MAP.wanted × ORIG_TABLE_NOS/ORIG_CAPTIONS
# (即 _write_tables_named 运行时真正取用的那两份), 产品改措辞 → 判据自动跟上。
def caption_hints() -> dict[str, list[str]]:
    """wanted/骨架表名 → 期望表题子串集 (与 _write_tables_named 的配对规则同构)

    配对规则 (与产品一致):
      - 槽位平行 (wanted/nos/caps 等长): wanted[i] ↔ 「nos[i] caps[i]」整行, 只认这一行;
      - 有取证但非平行 (2.1.1 拆表): 该节「表号 表题」全部行任一命中;
      - 无取证节 / nos≠caps 长度 (卫生异常, 由 hygiene 判红): 期望 = 内部名
        (产品回退分支的表题原文)。
    ⚠ 有取证时**不放内部名兜底** —— 否则「表题回退内部名」这一原始 bug 形态
    会自证清白 (内部名出现在表题里 = 判据命中 = 假绿)。"""
    from web.orig_table_map import ORIG_CAPTIONS, ORIG_TABLE_NOS
    from web.word_export import _SUB_TABLE_MAP, _match_wanted
    hints: dict[str, list[str]] = {}
    for sub, m in sorted(_SUB_TABLE_MAP.items()):
        if not m:
            continue
        wanted = m[1] or []
        nos = ORIG_TABLE_NOS.get(sub) or []
        caps = ORIG_CAPTIONS.get(sub) or []
        lines = [f"{n} {c}" for n, c in zip(nos, caps)] if nos and len(nos) == len(caps) else []
        parallel = bool(wanted) and bool(lines) and len(lines) == len(wanted)
        for w in wanted:
            hs = hints.setdefault(w, [])
            if parallel:
                i = _match_wanted(wanted, w)
                cand = [lines[i]] if i is not None else lines
            elif lines:
                cand = lines
            else:
                cand = [w]
            for c in cand:
                if c not in hs:
                    hs.append(c)
    return hints


CAP_HINT = caption_hints()


def extract_table_captions(doc) -> list[str]:
    """抓表题 (形如 '表3.1-1 xxx') — 表题通常在表前一段"""
    caps = []
    for p in doc.paragraphs:
        t = p.text.strip()
        m = re.match(r"^表\s*([\d.]+[-—]\d+)\s*(.*)$", t)
        if m:
            caps.append(t[:60])
    return caps


def table_shapes(doc) -> list[tuple[int, int]]:
    return [(len(t.rows), len(t.columns)) for t in doc.tables]


def main():
    pid = sys.argv[1] if len(sys.argv) > 1 else "c8c7ff0a4d"
    conn = connect_ro()
    conn.row_factory = sqlite3.Row
    data = json.loads(conn.execute("SELECT data FROM project WHERE id=?", (pid,)).fetchone()["data"])

    from web.derived_fields import derive_fields
    from web.table_skeleton import build_skeletons
    from web.external_tables import fill_one
    from web.word_export import export_docx

    data.update({k: v for k, v in derive_fields(data).items() if v})

    # ⚠ 与生产同构: assess 先算, build_skeletons 必须**带 assess** 才能填评估驱动表;
    #   有存储的 built_tables 就优先用它 (生产走的是 _get_assess 回填后的那份)。
    try:
        from knowledge.project_assess import assess_project
        from knowledge.oel import connect as oc
        _proj = {"id": pid, "data": data, "name": data.get("name") or "test",
                 "built_tables": data.get("built_tables", {}),
                 "hazard_grid": data.get("hazard_grid", []),
                 "emergency_supplies": data.get("emergency_supplies", []),
                 "location": data.get("location", ""), "company": data.get("company", "")}
        assess = assess_project(oc(), _proj)
        assess["_project_data"] = dict(_proj)
    except Exception as e:
        assess = {}
        print("assess err:", e)

    sk = dict(data.get("built_tables") or {})
    if not sk:
        sk = build_skeletons(data, assess or None)
    for n in list(sk):
        if isinstance(sk[n], dict) and sk[n].get("status") == "filling":
            t = fill_one(n, data)
            if t:
                sk[n] = t
    data["built_tables"] = sk
    if assess:
        assess["_project_data"] = dict(data)   # 让 _tables_for_sub 能取到骨架

    out = Path("/tmp/verify_export.docx")
    export_docx({"id": pid, "data": data, "name": data.get("name") or "test"}, assess, out)
    doc = docx.Document(str(out))

    caps = extract_table_captions(doc)
    shapes = table_shapes(doc)
    print("=" * 78)
    print(f"导出: {out}  ({out.stat().st_size} bytes)")
    print(f"  段落 {len(doc.paragraphs)} | 表格 {len(doc.tables)} | 表题 {len(caps)}")
    print()

    # 骨架里有行的表 (应该被导出)
    sk_with_rows = {n: len(t.get("rows") or []) for n, t in sk.items() if t.get("rows")}
    print(f"骨架表(有行) {len(sk_with_rows)} 张:")
    cap_blob = "\n".join(caps)
    missing = []
    for n, cnt in sorted(sk_with_rows.items()):
        hints = CAP_HINT.get(n, [n])
        found = any(h in cap_blob for h in hints)
        if not found:
            missing.append((n, cnt))
        print(f"  {'✓' if found else '✗'} {n:20} 骨架{cnt:3}行")
    print()
    print(f"骨架有行但导出未见表题的: {len(missing)}")
    for n, c in missing:
        print(f"   ✗ {n} ({c} 行)")

    print()
    print("=" * 78)
    print(f"导出 docx 里的表题 (共 {len(caps)}):")
    for c in caps:
        print("   ", c)
    print()
    print(f"表格形状 (行×列) 前 20:")
    for i, s in enumerate(shapes[:20], 1):
        print(f"   {i:2}. {s[0]}行 × {s[1]}列")


if __name__ == "__main__":
    main()

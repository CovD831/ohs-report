"""看真实新泰报告怎么写"识别出但未检测"的因素 / 岗位多因素表

用户原则: "关于规则或者是模板之类的, 看已有的报告", 不拍脑袋。
这是 Q2 的取证: 真实报告里"未检测"怎么表述。
"""
import re
from docx import Document

DOC = "/tmp/xt_probe/6-1新泰--预评（备案稿）.docx"
doc = Document(DOC)

paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
full = "\n".join(paras)

print("=" * 78)
print("① 含'未检测/未测/未进行检测'的段落")
print("=" * 78)
n = 0
for i, t in enumerate(paras):
    if re.search(r"未检测|未测|未进行检测|未列入检测", t):
        print(f"[{i}] {t[:200]}")
        n += 1
        if n >= 12:
            break
print(f"(命中 {n} 段)")

print()
print("=" * 78)
print("② 含'识别'与'检测'对比的表述")
print("=" * 78)
n = 0
for i, t in enumerate(paras):
    if re.search(r"识别(出|的).{0,20}(危害因素|职业病危害)", t) and len(t) < 300:
        print(f"[{i}] {t[:220]}")
        n += 1
        if n >= 8:
            break
print(f"(命中 {n} 段)")

print()
print("=" * 78)
print("③ 全部表格: 找'岗位/工种 × 因素'多因素表")
print("=" * 78)
for ti, tb in enumerate(doc.tables):
    try:
        hdr = [c.text.strip() for c in tb.rows[0].cells]
    except Exception:
        continue
    hs = " | ".join(hdr)
    if re.search(r"工种|岗位|接触", hs):
        print(f"\n--- 表{ti}  表头: {hs}  ({len(tb.rows)}行) ---")
        for r in tb.rows[1:6]:
            cells = [c.text.strip() for c in r.cells]
            print("   ", " | ".join(cells)[:190])
        if len(tb.rows) > 6:
            print(f"    ... 共 {len(tb.rows)} 行")

"""对 E2E 下载的报告跑全部检查器 (最终验收)

⚠ 这是唯一"用户真正拿到的东西" —— 不是内部函数产物, 是 HTTP 下载回来的 docx。
所有前面的验证都可能与真实路径有偏差, 这一步才是终局。
"""
import json
import re
import sqlite3
import sys
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

sys.path.insert(0, "/Users/abaaba/Projects/ohs-report")

DOC = Path(sys.argv[1] if len(sys.argv) > 1 else "/tmp/e2e_report_c64c1704bf.docx")
PID = sys.argv[2] if len(sys.argv) > 2 else "c64c1704bf"

doc = docx.Document(str(DOC))
print("=" * 80)
print(f"验收对象: {DOC.name} ({DOC.stat().st_size//1024} KB)")
print("=" * 80)

# ---- 1. 结构 ----
paras = [p.text.strip() for p in doc.paragraphs]
allt = "\n".join(paras)
heads = [t for t in paras if re.match(r"^\d{1,2}(\.\d{1,2}){0,3}\s{2,}\S", t) and len(t) <= 46]
caps = [t for t in paras if re.match(r"^表\s*[\d.]+[-—]", t)]
print(f"段落 {len(paras)} | 表格 {len(doc.tables)} | 表题 {len(caps)} | 标题 {len(heads)}")
nums = [re.match(r"^表\s*([\d.]+[-—]\d+)", c).group(1) for c in caps]
from collections import Counter
dup = [k for k, v in Counter(nums).items() if v > 1]
print(f"表题重号: {dup if dup else '无 ✓'}")
print(f"含'表格由系统'残留: {'有 ✗' if '表格由系统' in allt else '无 ✓'}")

# ---- 2. TOC 域 ----
xml = doc.element.body.xml
ins = re.findall(r"<w:instrText[^>]*>([^<]*)</w:instrText>", xml)
print(f"TOC 域: {[x.strip() for x in ins][:2] or '无 ✗'}")

# ---- 3. 字体 (标题/表题) ----
def font_of(r):
    if r is None or r._element.rPr is None:
        return (None, None)
    ft = r._element.rPr.find(qn("w:rFonts"))
    s = r._element.rPr.find(qn("w:sz"))
    return (ft.get(qn("w:eastAsia")) if ft is not None else None,
            s.get(qn("w:val")) if s is not None else None)

hf = Counter(); cf = Counter()
for p in doc.paragraphs:
    t = p.text.strip()
    if not t or len(t) > 46:
        continue
    if re.match(r"^\d{1,2}(\.\d{1,2}){0,3}\s{2,}\S", t) and p.runs:
        hf[font_of(p.runs[0])] += 1
    if re.match(r"^表\s*[\d.]+[-—]", t) and p.runs:
        cf[font_of(p.runs[0])] += 1
print(f"标题字体: {hf.most_common(2)}")
print(f"表题字体: {cf.most_common(2)}")

# ---- 4. 内容: 关键数据点 ----
print()
print("--- 关键数据点是否落到正文 ---")
import sqlite3
c = sqlite3.connect("/Users/abaaba/Projects/ohs-report/data/ohs.db")
c.row_factory = sqlite3.Row
row = c.execute("SELECT data FROM project WHERE id=?", (PID,)).fetchone()
if row:
    d = json.loads(row["data"])
    checks = [
        ("建设单位", d.get("company")),
        ("法定代表人", d.get("legal_rep")),
        ("注册资本", d.get("registered_capital")),
        ("成立日期", d.get("founded")),
        ("游离二氧化硅 3.39", "3.39"),
    ]
    for label, val in checks:
        if not val:
            print(f"  {label:16} 数据源缺失 (None)")
            continue
        hit = str(val) in allt
        print(f"  {label:16} {'✓ 正文含' if hit else '✗ 正文未出现'}  {str(val)[:34]}")

# ---- 5. 表格完整性 ----
gb = d.get("built_tables") or {} if row else {}
gb_rows = {n: len(t.get("rows") or []) for n, t in gb.items() if t.get("rows")}
print()
print(f"骨架有行表 {len(gb_rows)} 张 vs 导出 {len(doc.tables)} 张")

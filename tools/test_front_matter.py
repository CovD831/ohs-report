#!/usr/bin/env python3
"""前置页 (封面/声明页) 回归守卫 — 复刻真实备案稿格式

格式取证来源: 3 份真实报告 (长兴/浦发/新泰) span 级字号+坐标 + GBZ/T 196—2025 附录D
断言:
  ① 封面 p1: 企业名22粗 → 项目名18粗 → 报告书22粗 → 备案稿22粗 → 编号16粗 (全居中)
  ② 封面 p2: 机构名18粗 + 汉字年月16粗 (居中)
  ③ 声明页: '声  明'14粗居中 → 正文14pt 首行缩进2字 → 评价机构名称右对齐 → 法人代表缩进11字
  ④ 人员表: 4行 (项目负责人/报告书编写人/报告书审核人/报告书签发人), 无边框
  ⑤ 页码隔离: 前置(含封面/声明)无页码; 正文节有 PAGE 域
  ⑥ 编号生成: Y{年}-{3位序号} 递增; 已有编号幂等复用

无 DB / 无 LLM / 无外部依赖 — 纯 python-docx 结构断言
用法: .venv/bin/python tools/test_front_matter.py
"""
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.oxml.ns import qn

from web.front_matter import add_front_matter, gen_report_no, cn_year_month

OUT = ROOT / "data" / "_test_front_matter.docx"

FAILS = []


def check(name, cond, detail=""):
    mark = "✓" if cond else "✗"
    print(f"  {mark} {name}" + (f"  [{detail}]" if detail and not cond else ""))
    if not cond:
        FAILS.append(name)


def _sz(run):
    return run.font.size.pt if (run and run.font.size) else None


def main():
    # ── 构造 doc: 前置页 + 分节 + 占位正文 (模拟 word_export 接线) ──
    doc = Document()
    data = {
        "company": "长兴合成树脂（常熟）有限公司",
        "name": "年产12000吨不饱和聚酯树脂扩建项目",
        "org_name": "",  # 空 → 用默认机构名
        "report_no": "Y2026-001",
    }
    add_front_matter(doc, data)
    from docx.enum.section import WD_SECTION
    # fm6: 三节结构 (前置 / 目录·upperRoman / 正文·decimal)
    from web.word_export import (_setup_hf_styles, _attach_toc_hf, _attach_body_hf,
                                 _add_end_bookmark, _set_pg_num_type, add_toc)
    doc.settings.odd_and_even_pages_header_footer = True
    _setup_hf_styles(doc)
    _hf_title = f"{data['company']}{data['name']}职业病危害预评价报告书"
    _org = "江苏宁大卫防检测技术有限公司常熟分公司"
    toc_sec = doc.add_section(WD_SECTION.NEW_PAGE)
    _attach_toc_hf(toc_sec, doc, _hf_title, data["report_no"], _org)
    _set_pg_num_type(toc_sec, fmt="upperRoman", start=1)
    add_toc(doc)
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    _attach_body_hf(sec, doc, _hf_title, data["report_no"], _org)
    _set_pg_num_type(sec, fmt="decimal", start=1)
    doc.add_paragraph("正文占位")
    _add_end_bookmark(doc)
    doc.save(str(OUT))
    doc = Document(str(OUT))

    paras = [p for p in doc.paragraphs]
    nonempty = [p for p in paras if p.text.strip()]

    print("① 封面 p1 字段与字号")
    t0 = nonempty[0].text
    check("企业名 22pt粗居中", t0 == "长兴合成树脂（常熟）有限公司" and _sz(nonempty[0].runs[0]) == 22.0
          and nonempty[0].runs[0].font.bold, t0)
    check("项目名 18pt粗居中", nonempty[1].text.startswith("年产12000吨") and _sz(nonempty[1].runs[0]) == 18.0,
          nonempty[1].text)
    check("报告书 22pt粗", nonempty[2].text == "职业病危害预评价报告书" and _sz(nonempty[2].runs[0]) == 22.0)
    check("备案稿 22pt粗", nonempty[3].text == "（备案稿）" and _sz(nonempty[3].runs[0]) == 22.0)
    check("编号 16pt粗 (含 Y2026-001)", nonempty[4].text == "报告书编号：Y2026-001"
          and _sz(nonempty[4].runs[0]) == 16.0, nonempty[4].text)

    print("② 封面 p2 机构名+日期")
    org_paras = [p for p in nonempty if "宁大卫防" in p.text and p.alignment is not None]
    org_p = [p for p in nonempty if p.text.startswith("江苏宁大卫防检测技术有限公司")]
    check("机构名 18pt粗 (默认机构)", bool(org_p) and _sz(org_p[0].runs[0]) == 18.0,
          org_p[0].text if org_p else "缺失")
    dt_p = [p for p in nonempty if "年" in p.text and "月" in p.text
            and all(c in "〇一二三四五六七八九十" for c in p.text.replace("年", "").replace("月", ""))]
    check("汉字年月 16pt粗", bool(dt_p) and _sz(dt_p[0].runs[0]) == 16.0,
          dt_p[0].text if dt_p else "缺失")

    print("③ 声明页")
    decl = [p for p in nonempty if p.text.strip() == "声  明"]
    check("声明标题 14pt粗居中", bool(decl) and _sz(decl[0].runs[0]) == 14.0, "")
    body = [p for p in nonempty if "承担法律责任" in p.text]
    check("声明正文含 机构名+企业名+项目名", bool(body) and "长兴合成树脂" in body[0].text
          and "年产12000吨" in body[0].text and "承担法律责任" in body[0].text)
    sign = [p for p in nonempty if p.text.startswith("评价机构名称")]
    check("评价机构名称 右对齐", bool(sign) and str(sign[0].alignment) == "RIGHT (2)",
          str(sign[0].alignment) if sign else "缺失")
    rep = [p for p in nonempty if p.text.strip() == "法人代表："]
    ind_ok = False
    if rep:
        ind = rep[0]._p.find(qn("w:pPr") + "/" + qn("w:ind"))
        ind_ok = ind is not None and ind.get(qn("w:firstLineChars")) == "1100"
    check("法人代表 首行缩进11字", ind_ok)

    print("④ 人员表")
    ppl = doc.tables[0]
    labels = [ppl.rows[i].cells[0].text for i in range(len(ppl.rows))]
    check("4行标签齐全", labels == ["项目负责人：", "报告书编写人：", "报告书审核人：", "报告书签发人："],
          str(labels))
    borders = ppl._tbl.tblPr.find(qn("w:tblBorders"))
    allnone = borders is not None and all(c.get(qn("w:val")) == "none" for c in borders)
    check("无边框 (复刻浦发)", allnone)

    print("⑤ 页码隔离 (fm6: 三节)")
    s0 = doc.sections[0]
    instrs0 = s0.footer.paragraphs[0]._p.findall(".//" + qn("w:instrText")) if s0.footer.paragraphs else []
    check("前置节无页码域", not instrs0)
    check("节数=3 (前置/目录/正文)", len(doc.sections) == 3, str(len(doc.sections)))
    s1 = doc.sections[1]
    instrs1 = []
    for pp in s1.footer.paragraphs:
        instrs1 += [x.text for x in pp._p.findall(".//" + qn("w:instrText"))]
    check("目录节有 PAGE 域", any("PAGE" in (t or "") for t in instrs1), str(instrs1))
    s2 = doc.sections[2]
    instrs2 = []
    for pp in s2.footer.paragraphs:
        instrs2 += [x.text for x in pp._p.findall(".//" + qn("w:instrText"))]
    check("正文节有 PAGE 域", any("PAGE" in (t or "") for t in instrs2), str(instrs2))

    print("⑥ 编号生成")
    check("格式 Y{年}-{3位}", gen_report_no([], dt=__import__("datetime").datetime(2026, 5, 1)) == "Y2026-001")
    check("递增 (已有 007)", gen_report_no(["Y2026-007"], dt=__import__("datetime").datetime(2026, 5, 1)) == "Y2026-008")
    check("跨年隔离 (2025 的 020 不影响 2026)",
          gen_report_no(["Y2025-020"], dt=__import__("datetime").datetime(2026, 5, 1)) == "Y2026-001")
    check("中文年月", cn_year_month(__import__("datetime").datetime(2023, 7, 15)) == "二〇二三年七月")

    print("⑦ 页眉 (fm5)")
    s1 = doc.sections[2]           # 正文节 (fm6 后为第 3 节)
    hdr = s1.header
    htxt = hdr.paragraphs[0].text if hdr.paragraphs else ""
    check("页眉=企业名+项目名+报告书", htxt.startswith("长兴合成树脂") and "职业病危害预评价报告书" in htxt
          and "Y2026-001" in htxt, htxt)
    _ntabs = len(list(hdr.paragraphs[0]._p.iter(qn("w:tab")))) if hdr.paragraphs else 0
    check("页眉含 tab (编号右对齐用)", _ntabs >= 1, f"tabs={_ntabs}")
    _hxml = hdr.paragraphs[0]._p.xml if hdr.paragraphs else ""
    check("页眉编号含 noBreakHyphen", "noBreakHyphen" in _hxml, "")
    _hstyle = doc.styles["Header"]
    _hpPr = _hstyle.element.find(qn("w:pPr"))
    _hbdr = _hpPr.find(qn("w:pBdr")) if _hpPr is not None else None
    check("Header 样式含下边框", _hbdr is not None and _hbdr.find(qn("w:bottom")) is not None)
    _htabs = _hpPr.find(qn("w:tabs")) if _hpPr is not None else None
    _clears = len([t for t in _htabs if t.get(qn("w:val")) == "clear"]) if _htabs is not None else 0
    _right = [t.get(qn("w:pos")) for t in _htabs if t.get(qn("w:val")) == "right"] if _htabs is not None else []
    check("Header/Footer 样式 clear 模板tab + right@8334", _clears == 2 and _right == ["8334"],
          f"clears={_clears} right={_right}")

    print("⑧ 页脚奇偶镜像 (fm5)")
    s2 = doc.sections[2]
    f_odd = s2.footer.paragraphs[0].text
    f_even = s2.even_page_footer.paragraphs[0].text
    check("奇数页 机构名在前", f_odd.startswith("江苏宁大卫防"), f_odd)
    check("偶数页 页码在前 (镜像)", f_even.startswith("第"), f_even)
    check("奇数页含 第X页共Y页", "第" in f_odd and "页 共" in f_odd and "页" in f_odd, f_odd)
    o_instr = []
    for pp in s2.footer.paragraphs:
        o_instr += [x.text for x in pp._p.findall(".//" + qn("w:instrText"))]
    check("合计页数用 PAGEREF (非 NUMPAGES/SECTIONPAGES)",
          any("PAGEREF _ohs_report_end" in (t or "") for t in o_instr)
          and not any("NUMPAGES" in (t or "") or "SECTIONPAGES" in (t or "") for t in o_instr), str(o_instr))
    check("evenAndOddHeaders 开启", doc.settings.odd_and_even_pages_header_footer is True)
    _bx = doc.element.body.findall(".//" + qn("w:bookmarkStart"))
    check("文末书签存在", any(b.get(qn("w:name")) == "_ohs_report_end" for b in _bx))

    print("⑨ 目录节罗马页码 (fm6)")
    stoc = doc.sections[1]
    pg = stoc._sectPr.find(qn("w:pgNumType"))
    check("目录节 pgNumType fmt=upperRoman",
          pg is not None and pg.get(qn("w:fmt")) == "upperRoman",
          str(dict((k.split('}')[-1], v) for k, v in (pg.attrib.items() if pg is not None else []))))
    check("目录节页码从 I 起 (start=1)",
          pg is not None and pg.get(qn("w:start")) == "1",
          str(pg.get(qn("w:start")) if pg is not None else None))
    tf_odd = stoc.footer.paragraphs[0].text
    tf_even = stoc.even_page_footer.paragraphs[0].text
    check("目录页脚=机构名+裸罗马数字 (奇数页机构名在前)", tf_odd.startswith("江苏宁大卫防")
          and "第" not in tf_odd and "页" not in tf_odd, tf_odd)
    check("目录页脚 偶数页镜像 (罗马数字在前)", "第" not in tf_even and "江苏宁大卫防" in tf_even
          and tf_even.index("江苏宁大卫防") > 0, tf_even)
    tobj = []
    for pp in stoc.footer.paragraphs:
        tobj += [x.text for x in pp._p.findall(".//" + qn("w:instrText"))]
    check("目录页脚仅 PAGE 域 (无 PAGEREF 共Y页)",
          any("PAGE" in (t or "") for t in tobj) and not any("PAGEREF" in (t or "") for t in tobj), str(tobj))
    _sh = stoc.header.paragraphs[0].text if stoc.header.paragraphs else ""
    check("目录节页眉同正文 (标题+编号)", _sh.startswith("长兴合成树脂") and "Y2026-001" in _sh, _sh)
    _spg = doc.sections[2]._sectPr.find(qn("w:pgNumType"))
    check("正文节 pgNumType decimal start=1",
          _spg is not None and _spg.get(qn("w:fmt")) == "decimal" and _spg.get(qn("w:start")) == "1",
          str(dict((k.split('}')[-1], v) for k, v in (_spg.attrib.items() if _spg is not None else []))))

    OUT.unlink(missing_ok=True)

    print()
    if FAILS:
        print(f"前置页守卫: ✗ 失败 {len(FAILS)} 项: {FAILS}")
        return 1
    print("前置页守卫: ✓ 全部通过")
    return 0


if __name__ == "__main__":
    sys.exit(main())

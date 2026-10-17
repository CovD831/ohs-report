"""回归测试: 表挂载点必须可达 + 骨架表必须在导出里出现

防的是这类静默失败:
  _SUB_TABLE_MAP 把表挂到一个**结构里不存在的节号** → 导出 walk 永远不到
  → 整张表消失, 且没有任何报错 (2026-09: 工艺检查表挂 3.5.3, 而 3.5 无子节)

用法: python tools/test_table_mounts.py
退出码 0 = 通过
"""
import sys
from pathlib import Path

ROOT = next((p for p in (Path("/app"), Path(__file__).resolve().parent.parent)
             if (p / "web" / "app.py").exists()), Path(__file__).resolve().parent.parent)
sys.path.insert(0, str(ROOT))


def reachable_mounts() -> tuple[list[tuple[str, list[str]]], set[str]]:
    from tools.check_mount_points import all_sections
    from web.word_export import _SUB_TABLE_MAP
    valid = all_sections()
    bad = []
    for sub, m in sorted(_SUB_TABLE_MAP.items()):
        if not m:
            continue
        if sub not in valid:
            bad.append((sub, m[1]))
    return bad, valid


def skeleton_names() -> set[str]:
    from web.word_export import _SUB_TABLE_MAP
    out = set()
    for sub, m in sorted(_SUB_TABLE_MAP.items()):
        if m:
            out.update(m[1])
    return out


def test_mounts_reachable() -> int:
    """挂载点必须在报告结构里存在 (允许已知例外, 但必须显式登记)"""
    # 已知例外: 5.1.1.2 由产品单元展开而来, 非 SUBS 字面节号 → 可接受
    ALLOW = {"5.1.1.2"}
    bad, valid = reachable_mounts()
    hard = [(s, n) for s, n in bad if s not in ALLOW]
    print(f"挂载点检查: 结构 {len(valid)} 节, 不可达 {len(bad)} 个 "
          f"(允许例外 {len(bad) - len(hard)})")
    for s, names in bad:
        tag = "允许" if s in ALLOW else "★错误"
        print(f"  [{tag}] {s} -> {names}")
    if hard:
        print(f"  ✗ {len(hard)} 个挂载点不可达且未登记例外")
        return len(hard)
    print("  ✓ 通过")
    return 0


def test_orig_map_hygiene() -> int:
    """ORIG 编号/表题映射卫生 (防章号平移后死键/槽位不平行 → 表题静默回退内部名)

    2026-09 实例: 挂载 3.5.3→3.5 / 5.4.3→10.1 / 10.2→11.2 迁移后映射留旧键,
    表题静默回退内部名 (「本项目工艺检查及评价」→「工艺检查表」)。
    规则: ① 键必须仍是挂载节 (无死键); ② 有映射的节编号/表题列表平行,
    且与 wanted 名单位次一一对应 (2.1.1 应急三拆登记例外)。
    """
    from web.orig_table_map import ORIG_CAPTIONS, ORIG_TABLE_NOS
    from web.word_export import _SUB_TABLE_MAP

    mounts = {k for k, v in _SUB_TABLE_MAP.items() if v}
    split_ok = {"2.1.1"}  # wanted 1 → 编号/表题 3 (应急药品/物资/洗眼器 三拆)
    bad = 0
    dead = sorted((set(ORIG_TABLE_NOS) | set(ORIG_CAPTIONS)) - mounts)
    if dead:
        print(f"  ✗ ORIG 映射死键 (挂载已迁走): {dead}")
        bad += len(dead)
    for k in sorted(set(ORIG_TABLE_NOS) | set(ORIG_CAPTIONS)):
        nos = ORIG_TABLE_NOS.get(k) or []
        caps = ORIG_CAPTIONS.get(k) or []
        if len(nos) != len(caps):
            print(f"  ✗ {k}: 编号 {len(nos)} ≠ 表题 {len(caps)} (须平行)")
            bad += 1
            continue
        wanted = (_SUB_TABLE_MAP.get(k) or (None, []))[1]
        if k not in split_ok and wanted and len(wanted) != len(nos):
            print(f"  ✗ {k}: wanted {len(wanted)} ≠ 槽位 {len(nos)} (按名配对失效)")
            bad += 1
    tag = "✗" if bad else "✓"
    print(f"  {tag} ORIG 映射: {len(mounts)} 挂载 / {len(ORIG_TABLE_NOS)} 编号 / {len(ORIG_CAPTIONS)} 表题"
          + (f" — {bad} 项异常" if bad else " — 无死键, 槽位平行"))
    return bad


def test_tables_export(tmpdir=None) -> int:
    """真项目跑导出, 骨架里有行的表必须都在 docx 里出现"""
    import json
    import re
    import sqlite3

    import docx

    # 只读工厂 (S1c″ D4′): 原来这里写死 `ROOT/"data"/"ohs.db"` + 可写 sqlite3.connect,
    # 于是 ①认不到 OHS_DATA_DIR (临时卷演练测不到它) ②容器里就是挂进去的生产库
    # ③缺库时 print 一句 return 0 = "绿而没测"。现在三条一起修。
    from knowledge.oel import connect_ro
    from knowledge.paths import DB_PATH
    try:
        conn = connect_ro()
    except Exception as e:
        print(f"  ✗ 官方数据层不可读, 导出检查无法执行: {type(e).__name__}: {e} [DB={DB_PATH}]"
              f" —— 记 1 项失败, 不许 print 一句就走")
        return 1
    conn.row_factory = sqlite3.Row
    # 优先 d4c2b5450c (服务器真项目), 否则挑任一有数据的项目
    row = conn.execute("SELECT id, data FROM project WHERE id=?",
                       ("d4c2b5450c",)).fetchone()
    if not row:
        for r in conn.execute("SELECT id, data FROM project"):
            try:
                dd = json.loads(r["data"])
            except Exception:
                continue
            if dd.get("built_tables") or dd.get("equipment_detail") or dd.get("equipment"):
                row = r
                break
    if not row:
        print(f"  ✗ 库里没有任何带数据的项目, 导出检查没有执行对象 [DB={DB_PATH}]"
              f" —— 记 1 项失败 (旧版 print 一句 return 0 = 绿而没测)")
        return 1
    pid, data = row["id"], json.loads(row["data"])
    conn.close()

    from web.derived_fields import derive_fields
    from web.table_skeleton import build_skeletons
    from web.external_tables import fill_one
    from web.word_export import export_docx

    data.update({k: v for k, v in derive_fields(data).items() if v})
    sk = build_skeletons(data)
    for n in list(sk):
        if sk[n].get("status") == "filling":
            t = fill_one(n, data)
            if t:
                sk[n] = t
    data["built_tables"] = sk
    try:
        from knowledge.project_assess import assess_project
        # 单独再开一次只读连接: assess_project 会在收尾时 conn.close() (project_assess:694),
        # 与上面取项目数据的那条不能复用。旧版这里用 `from knowledge.oel import connect as oc`
        # + `assess_project(oc(), data)` —— 写侧工厂 (mkdir + executescript(DDL)),
        # 且 `except Exception: assess = {}` 把它的失效整个吞掉。
        assess = assess_project(connect_ro(), data)
        assess_failed = 0
    except Exception as e:
        print(f"  ✗ assess_project 抛错 → assess 退化成空表, 下面的挂载检查不可信: "
              f"{type(e).__name__}: {e}")
        assess, assess_failed = {}, 1

    out = Path(tmpdir or "/tmp") / "test_mount_export.docx"
    export_docx({"id": pid, "data": data, "name": "t"}, assess, out)
    doc = docx.Document(str(out))
    caps = [p.text.strip() for p in doc.paragraphs
            if re.match(r"^表\s*[\d.]+[-—]\d+", p.text.strip())]
    blob = "\n".join(caps)
    with_rows = {n for n, t in sk.items() if t.get("rows")}

    sys.path.insert(0, str(ROOT / "tools"))
    # 期望表题**从产品真源推导** (_SUB_TABLE_MAP × ORIG_TABLE_NOS/ORIG_CAPTIONS,
    # 见 verify_export_tables.caption_hints) —— 旧版这里 import 手抄别名表 CAP_HINT,
    # 2026-10 「设备设施」旧措辞漂移 → 表明明导出却被判缺失 (判据错, 不是产品缺陷)。
    # 必须在导出后调用: 变异模式 (run_mutations) monkeypatch 的就是这两份真源。
    from verify_export_tables import caption_hints
    hints = caption_hints()
    missing = [n for n in with_rows
               if not any(h in blob for h in hints.get(n, [n]))]

    # 跨通道判据 (变异自检 M1/M2 实测补强): sk/with_rows 的名单派生自 _SUB_TABLE_MAP
    # 本身 —— 「删挂载点 / 把表名写错」时名单跟着缩, 只对照 sk 等于拿被测对象的尺子
    # 量它自己 (2026-09 工艺检查表整表消失正是这条漏抓路径)。
    # web.table_builder.build_data_tables 是**数据层独立产地** (不消费 _SUB_TABLE_MAP):
    # 它从项目原始数据真造出行来的表, 必须出现在导出里 —— 名单不随挂载变异缩, 无循环。
    from web.table_builder import build_data_tables
    for _n, _t in build_data_tables(data).items():
        if isinstance(_t, dict) and _t.get("rows") and _n not in with_rows:
            if not any(h in blob for h in hints.get(_n, [_n])):
                missing.append(_n)
    missing = sorted(set(missing))
    print(f"导出检查: 项目 {pid} | 骨架有行 {len(with_rows)} | docx 表 {len(doc.tables)} "
          f"| 表题 {len(caps)}")
    if missing:
        print(f"  ✗ 骨架有行但导出未见: {missing}")
        return len(missing) + assess_failed
    print(f"  {'✓' if not assess_failed else '✗'} {len(with_rows)}/{len(with_rows)} 全部导出"
          + ("" if not assess_failed else " (但 assess 为空, 见上面那行 ✗)"))
    return assess_failed


def run_mutations() -> int:
    """判据自证 (「当证据的检查先证明能红」): 向两份真源注入破坏, 每条**必须**每次都红。

    全程 in-process monkeypatch + finally 还原, 不写任何源文件; 导出只落 /tmp。
    变异对象就是判据与被测产品共同读的那两份真源 (_SUB_TABLE_MAP /
    ORIG_TABLE_NOS+ORIG_CAPTIONS), 所以不存在「判据与产品各看一份名单」的假绿通道。
    返回没被抓住的变异数 (0 = 判据有牙齿)。
    """
    from web import orig_table_map, word_export

    orig_map = word_export._SUB_TABLE_MAP
    orig_nos = orig_table_map.ORIG_TABLE_NOS
    orig_caps = orig_table_map.ORIG_CAPTIONS
    escaped = 0

    def verdict(label, channel, fn, expect_pos):
        nonlocal escaped
        got = fn()
        ok = (got > 0) == expect_pos
        print(f"  [{'✓ 抓到' if ok else '✗ 漏抓!!'}] {label} [{channel}] → 红 {got} 项")
        if not ok:
            escaped += 1

    print("变异自检 (每条都必须红; 时有时无的判据不算判据):")
    try:
        # M1 删挂载点: 摘掉 _SUB_TABLE_MAP["3.6.2"] (设备明细表的唯一挂载)。
        #   必红机理: ① 挂载没了 → _tables_for_sub 无映射 → 246 行的表整张不上纸,
        #   且 caption_hints 也从同一份 map 推导, hints 退化为 [骨架名], 而真表题
        #   不含骨架名 → export 判据必红; ② ORIG["3.6.2"] 成死键 → hygiene 必红。
        word_export._SUB_TABLE_MAP = {k: v for k, v in orig_map.items() if k != "3.6.2"}
        verdict("M1 删挂载点(3.6.2)", "hygiene", test_orig_map_hygiene, True)
        verdict("M1 删挂载点(3.6.2)", "export", test_tables_export, True)
    finally:
        word_export._SUB_TABLE_MAP = orig_map
    print()
    try:
        # M2 把表名写错: 11.5 挂载 wanted 改成不存在的「体检周期表」。
        #   必红机理: _tables_for_sub 按 wanted 精确过滤 → 体检项目与周期表被丢弃,
        #   不上纸; hygiene 查不到它 (长度仍平行) —— 正是 export 通道必须独立有牙齿的变异。
        word_export._SUB_TABLE_MAP = dict(orig_map)
        word_export._SUB_TABLE_MAP["11.5"] = ("11", ["体检周期表"])
        verdict("M2 表名写错(11.5)", "export", test_tables_export, True)
    finally:
        word_export._SUB_TABLE_MAP = orig_map
    print()
    try:
        # M3 编号/表题列表做成不平行: ORIG_CAPTIONS["5.3"] 砍成 2 条 (编号仍是 3 条)。
        #   必红机理: hygiene 的平行性检查是纯数据比对, 不依赖导出 → 每次都红。
        orig_table_map.ORIG_CAPTIONS = dict(orig_caps)
        orig_table_map.ORIG_CAPTIONS["5.3"] = orig_caps["5.3"][:2]
        verdict("M3 编号/表题不平行(5.3)", "hygiene", test_orig_map_hygiene, True)
    finally:
        orig_table_map.ORIG_CAPTIONS = orig_caps
    print()
    try:
        # M4 原始 bug 回归: 把 4.4 系列 ORIG 键改回编号前缀「4.4」(10-04 的错法)。
        #   必红机理: 挂载在叶子节 4.4.2 → 「4.4」成死键 (hygiene 红);
        #   槽位映射失效 → 表题回退内部名 + 编号跟随数据出现序 (export 的
        #   精确「表号 表题」期望行落空)。这条钉死本批修的那个 bug 本身。
        orig_table_map.ORIG_TABLE_NOS = dict(orig_nos)
        orig_table_map.ORIG_CAPTIONS = dict(orig_caps)
        orig_table_map.ORIG_TABLE_NOS["4.4"] = orig_table_map.ORIG_TABLE_NOS.pop("4.4.2")
        orig_table_map.ORIG_CAPTIONS["4.4"] = orig_table_map.ORIG_CAPTIONS.pop("4.4.2")
        verdict("M4 ORIG 键回退成 4.4", "hygiene", test_orig_map_hygiene, True)
    finally:
        orig_table_map.ORIG_TABLE_NOS = orig_nos
        orig_table_map.ORIG_CAPTIONS = orig_caps
    print()
    # 对照: 全部还原后 hygiene 必须回绿 (红不是永久红, 是变异引起的)
    word_export._SUB_TABLE_MAP = orig_map
    verdict("还原对照", "hygiene", test_orig_map_hygiene, False)
    return escaped


if __name__ == "__main__":
    bad = test_mounts_reachable()
    print()
    bad += test_orig_map_hygiene()
    print()
    bad += test_tables_export()
    print()
    if "--mutations" in sys.argv:
        bad += run_mutations()
        print()
    print("结果:", "✓ 通过" if bad == 0 else f"✗ 失败 {bad} 项")
    sys.exit(1 if bad else 0)

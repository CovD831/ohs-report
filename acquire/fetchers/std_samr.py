"""全国标准信息公共服务平台 — 标准目录/状态抓取器
已验证: JSON API 直连可用, 返回标准现行/废止状态 (引用时效核验权威源)
注意: 该平台只覆盖 GB/GB/T/行业/地方/团体/企业标准, 不含 GBZ 系列 (见 nhc_gbz.py)
输出: acquire/raw/std_samr/standards_status.json
"""
import json
import re
import sys
import urllib.parse
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))  # 项目根
from acquire.base import PoliteFetcher  # noqa: E402

BASE = "https://std.samr.gov.cn/gb/search/gbQueryPage"
RAW_DIR = Path(__file__).resolve().parent.parent / "raw" / "std_samr"

# 注: 仅 GB/GB/T 体系 (GBZ/T 196-2025 第10章引用的国标 + 行业分类)
NEEDED = [
    "GB/T 4754-2017",    # 国民经济行业分类 (风险分类前置)
    "GB 5083",           # 生产设备安全卫生设计总则
    "GB/T 12801",        # 生产过程安全卫生要求总则
    "GB 39800",          # 个体防护装备配备规范
    "GB/T 18664",        # 呼吸防护用品的选择、使用与维护
    "GB/T 23466",        # 护听器的选择指南
    "GB/T 38144",        # 眼面部防护
    "GB 50187",          # 工业企业总平面设计规范
    "GB 50019",          # 工业建筑供暖通风与空气调节设计规范
    "GB 50033",          # 建筑采光设计标准
    "GB/T 50034",        # 建筑照明设计标准
    "GB 50073",          # 洁净厂房设计规范
    "GB/T 16758",        # 排风罩的分类及技术条件
    "GB/T 50087",        # 工业企业噪声控制设计规范
]


def strip_tags(s):
    return re.sub(r"<[^>]+>", "", s or "").strip()


# 标准引用形态: GB 30871—2022 / GB/T 29639-2020 / GBZ2.1-2019
_RE_REF = re.compile(r"\b(GB(?:/T|/Z)?|GBZ(?:/T)?)\s*(\d+(?:\.\d+)?)\s*[—\-–−]\s*(\d{4})", re.I)


def needed_from_reports() -> list[str]:
    """从**报告产物 + fixed_texts** 里扫实际引用的国标, 推导采集需求 (**通用性**: 不写死清单)

    数据源 (取并集):
      1. web/fixed_texts_data.json —— 法律依据固定文本
      2. data/report_*.docx 产物正文 —— 真实生成稿引用了什么
    只收 GB/GB/T/GB/Z (GBZ 走 nhc_gbz.py; 我方采集器只补 SAMR 侧盲区)
    """
    import json
    from pathlib import Path as _P
    root = _P(__file__).resolve().parent.parent.parent
    found: set[str] = set()
    texts: list[str] = []
    # 1) fixed_texts
    ft = root / "web" / "fixed_texts_data.json"
    if ft.exists():
        try:
            texts.append(ft.read_text(encoding="utf-8"))
        except Exception:
            pass
    # 2) 产物 docx
    try:
        import docx  # type: ignore
        for f in sorted((root / "data").glob("report_*.docx")):
            try:
                d = docx.Document(str(f))
                texts.append("\n".join(p.text for p in d.paragraphs))
                for t in d.tables:
                    for r in t.rows:
                        texts.append(" ".join(c.text for c in r.cells))
            except Exception:
                continue
    except Exception:
        pass
    for txt in texts:
        for pre, num, year in _RE_REF.findall(txt or ""):
            p = pre.upper().replace(" ", "")
            # 只收 GB / GB/T / GB/Z (GBZ 由 nhc_gbz 负责)
            if p.startswith("GBZ"):
                continue
            found.add(f"{p} {num}-{year}")
    return sorted(found)


def fetch(standards: list[str] | None = None) -> dict:
    """按标准号查询, 落盘 JSON (温和抓取: 指数退避+限速由 PoliteFetcher 保证)

    standards 传 None 时: 用 **NEEDED ∪ 报告实际引用推导**(见 needed_from_reports),
    避免硬编码清单与实际引用脱节 (2026-10 教训: NEEDED 里没有 GB 30871/GB/T 29639,
    报告引了却查不到状态, 无法判定是否废止)。
    """
    if standards is None:
        standards = sorted(set(NEEDED) | set(needed_from_reports()))
    fetcher = PoliteFetcher("std_samr")
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    out = {}
    for code in (standards or NEEDED):
        try:
            url = f"{BASE}?searchText={urllib.parse.quote(code)}&type=all"
            data = fetcher.get_json(url)
            # 响应结构: 有时顶层直接带 total/rows, 有时包在 data 里 (实测定两态)
            payload = data.get("data") if isinstance(data.get("data"), dict) else data
            rows = []
            for r in payload.get("rows", []):
                rows.append({
                    "code": strip_tags(r.get("C_STD_CODE")),
                    "name": r.get("C_C_NAME"),
                    "state": r.get("STATE"),          # 现行 / 废止
                    "nature": r.get("STD_NATURE"),    # 推荐性/强制性
                    "issue_date": r.get("ISSUE_DATE"),
                    "act_date": r.get("ACT_DATE"),
                    "id": r.get("id"),                # 详情页 gbDetailed?id=
                })
            out[code] = rows
            print(f"OK  {code} -> {[(r['code'], r['name'], r['state']) for r in rows[:3]]}")
        except Exception as e:
            out[code] = []
            print(f"ERR {code}: {e}")
    (RAW_DIR / "standards_status.json").write_text(
        json.dumps(out, ensure_ascii=False, indent=2))
    return out


if __name__ == "__main__":
    fetch()
